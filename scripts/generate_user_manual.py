from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Manual_de_Usuario_TicketFlow.docx"
PURPLE = "6755F5"
NAVY = "172033"
LIGHT = "F1EFFF"
GRAY = "687083"
GREEN = "E9F9EF"
YELLOW = "FFF5D8"
RED = "FFEAED"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(9)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True, "FFFFFF")
        shade(table.rows[0].cells[i], PURPLE)
        if widths:
            table.rows[0].cells[i].width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(table.rows) % 2 == 1:
                shade(cells[i], "F8F8FC")
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_steps(doc, steps):
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(step)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_note(doc, title, text, kind="info"):
    fills = {"info": LIGHT, "success": GREEN, "warning": YELLOW, "danger": RED}
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fills[kind])
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title + ": ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.3)
sec.right_margin = Cm(2.3)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size, color in [("Title", 34, NAVY), ("Heading 1", 22, PURPLE), ("Heading 2", 15, NAVY), ("Heading 3", 12, PURPLE)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

header = sec.header.paragraphs[0]
header.text = "TicketFlow  |  Manual de usuario"
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
add_page_number(sec.footer.paragraphs[0])

# Portada
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
r = p.add_run("T")
r.bold = True
r.font.size = Pt(42)
r.font.color.rgb = RGBColor.from_string(PURPLE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TicketFlow")
r.bold = True
r.font.size = Pt(30)
r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph("SISTEMA DE VENTA Y GESTIÓN DE ENTRADAS")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
p.runs[0].bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(38)
r = p.add_run("MANUAL DE USUARIO")
r.bold = True
r.font.size = Pt(24)
p = doc.add_paragraph("Funciones principales, roles y procesos operativos")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(14)
p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(100)
p.add_run("Versión 1.0\n").bold = True
p.add_run(f"Fecha de elaboración: {date.today().strftime('%d/%m/%Y')}\n")
p.add_run("Aplicación desarrollada con CodeIgniter 4 y PHP")
doc.add_page_break()

add_heading(doc, "Control del documento", 1)
add_table(doc, ["Campo", "Información"], [
    ("Documento", "Manual de usuario de TicketFlow"),
    ("Versión", "1.0"),
    ("Sistema", "Venta de entradas para eventos"),
    ("Entorno documentado", "Aplicación web local: http://localhost:8080"),
    ("Público objetivo", "Administradores, organizadores, vendedores, operadores de acceso y clientes"),
])
add_note(doc, "Alcance", "Este documento describe las funciones disponibles actualmente. El pago web es simulado y no procesa tarjetas ni dinero real.")

add_heading(doc, "Índice", 1)
for item in [
    "1. Introducción y conceptos generales", "2. Acceso al sistema", "3. Funciones del Administrador",
    "4. Funciones del Organizador", "5. Funciones del Vendedor", "6. Funciones de Control de acceso",
    "7. Funciones del Cliente", "8. Flujos completos de venta", "9. Estados del sistema",
    "10. Solución de problemas", "11. Seguridad y buenas prácticas", "12. Glosario y listas de verificación",
]:
    doc.add_paragraph(item)
doc.add_page_break()

add_heading(doc, "1. Introducción y conceptos generales", 1)
doc.add_paragraph("TicketFlow es una aplicación web para administrar eventos, programar funciones, controlar aforo, configurar tipos de entrada y descuentos, registrar ventas, emitir boletos y validar el ingreso de asistentes.")
add_heading(doc, "1.1 Objetivos del sistema", 2)
add_bullets(doc, [
    "Publicar un catálogo de eventos y funciones disponibles.",
    "Gestionar precios, localidades, cupos y límites de compra.",
    "Permitir compras web simuladas, reservas y ventas presenciales.",
    "Generar entradas individuales con códigos únicos.",
    "Evitar que una misma entrada sea utilizada más de una vez.",
    "Separar las responsabilidades mediante usuarios y roles.",
])
add_heading(doc, "1.2 Roles", 2)
add_table(doc, ["Rol", "Responsabilidad principal", "Acceso principal"], [
    ("Administrador", "Supervisión completa del sistema", "Todos los módulos"),
    ("Organizador", "Configuración comercial y operativa de eventos", "Eventos, funciones, tipos y descuentos"),
    ("Vendedor", "Ventas presenciales y confirmación de reservas", "Punto de venta e historial"),
    ("Control de acceso", "Validación de boletos en el ingreso", "Control de acceso e historial"),
    ("Cliente", "Consulta, compra y reserva de entradas", "Catálogo y Mis entradas"),
])
add_note(doc, "Principio", "Una persona cliente utiliza una sola cuenta. Esa cuenta puede tener muchas compras web, reservas y ventas presenciales.", "success")

add_heading(doc, "2. Acceso al sistema", 1)
add_heading(doc, "2.1 Acceso del personal", 2)
doc.add_paragraph("Dirección: http://localhost:8080/login")
add_steps(doc, ["Abrir la dirección del panel.", "Escribir el correo institucional.", "Escribir la contraseña.", "Seleccionar Iniciar sesión.", "El sistema mostrará los módulos permitidos para el rol."])
add_heading(doc, "2.2 Acceso del cliente", 2)
doc.add_paragraph("Dirección: http://localhost:8080/ingresar")
add_steps(doc, ["Seleccionar Ingresar desde el catálogo.", "Escribir correo y contraseña.", "Confirmar el acceso.", "Utilizar Eventos para comprar o Mis entradas para consultar boletos y reservas."])
add_heading(doc, "2.3 Registro de clientes", 2)
add_steps(doc, ["Seleccionar Crear cuenta.", "Completar nombres, apellidos, correo y teléfono.", "Crear una contraseña segura y confirmarla.", "Enviar el formulario.", "La sesión del cliente se inicia automáticamente."])
add_note(doc, "Contraseña segura", "Debe contener al menos 10 caracteres, mayúscula, minúscula, número y símbolo.", "warning")

add_heading(doc, "3. Funciones del Administrador", 1)
add_heading(doc, "3.1 Panel de resumen", 2)
add_bullets(doc, ["Revisar eventos activos.", "Consultar entradas vendidas, aforo disponible e ingresos.", "Consultar próximas funciones y su ocupación.", "Acceder rápidamente a la configuración principal."])
add_heading(doc, "3.2 Gestión de eventos", 2)
add_steps(doc, ["Abrir Eventos.", "Seleccionar Nuevo evento.", "Asignar organizador, nombre, descripción, categoría e imagen.", "Guardar inicialmente como borrador.", "Agregar funciones, tipos de entrada y cupos.", "Publicar cuando la configuración esté completa."])
add_note(doc, "Publicación", "Un evento publicado aparece en el catálogo público. La eliminación utilizada por el sistema es lógica, por lo que el registro no debe borrarse directamente desde la base.", "warning")
add_heading(doc, "3.3 Funciones", 2)
add_bullets(doc, ["Definir nombre de la función, fecha de inicio y finalización.", "Configurar inicio y fin de la venta pública.", "Registrar recinto, dirección, ciudad y aforo total.", "Cambiar el estado entre programada, en curso, cancelada o finalizada."])
add_heading(doc, "3.4 Tipos de entrada", 2)
add_steps(doc, ["Abrir Tipos de entrada o ingresar desde una función.", "Seleccionar General, VIP u otra denominación.", "Indicar precio y cupo.", "Definir el límite máximo por compra.", "Guardar y verificar que la suma de cupos no supere el aforo."])
add_heading(doc, "3.5 Descuentos", 2)
add_bullets(doc, ["Crear códigos únicos.", "Aplicar un valor fijo o un porcentaje.", "Asociar el código a un evento o a todos.", "Configurar vigencia, compra mínima y límite de usos.", "Activar o desactivar la promoción."])
add_heading(doc, "3.6 Usuarios y roles", 2)
add_steps(doc, ["Abrir Usuarios y roles.", "Crear o editar una cuenta.", "Asignar Administrador, Organizador, Vendedor, Control de acceso o Cliente.", "Activar, bloquear o desbloquear la cuenta.", "Cambiar la contraseña cuando corresponda."])
add_note(doc, "Protección", "El administrador no puede quitarse su propio rol ni bloquearse a sí mismo. Esta medida evita dejar el sistema sin acceso administrativo.", "info")

add_heading(doc, "4. Funciones del Organizador", 1)
doc.add_paragraph("El Organizador configura la oferta del evento. No administra usuarios ni valida boletos en la puerta.")
add_bullets(doc, [
    "Crear y editar eventos asignados.", "Programar funciones y fechas de venta.", "Definir recinto y aforo.",
    "Crear tipos de entrada y controlar que los cupos no excedan el aforo.", "Crear descuentos y controlar su vigencia.",
])
add_heading(doc, "4.1 Orden recomendado de configuración", 2)
add_steps(doc, ["Crear el evento.", "Crear al menos una función.", "Configurar tipos de entrada y cupos.", "Crear descuentos si son necesarios.", "Revisar fechas y aforo.", "Publicar el evento."])

add_heading(doc, "5. Funciones del Vendedor", 1)
doc.add_paragraph("El Vendedor trabaja en ventanilla. Registra ventas presenciales, consulta compras web y confirma el pago de reservas pendientes. No valida el ingreso al evento.")
add_heading(doc, "5.1 Venta presencial", 2)
add_steps(doc, [
    "Abrir Punto de venta.", "Escribir el correo del cliente.", "Si la cuenta existe, verificar el nombre recuperado automáticamente.",
    "Si el correo es nuevo, completar nombre y apellido.", "Seleccionar evento, función y tipo de entrada.",
    "Indicar cantidad y método recibido.", "Aplicar un código promocional válido si corresponde.",
    "Recibir el dinero.", "Seleccionar Vender y emitir.", "Consultar o imprimir las entradas generadas.",
])
add_note(doc, "Importante", "El campo promocional recibe códigos como VERANO20. Los códigos ENT- pertenecen a entradas y no deben escribirse como descuento.", "warning")
add_heading(doc, "5.2 Confirmar una reserva pagada en ventanilla", 2)
add_steps(doc, [
    "Abrir Historial de ventas.", "Localizar el código RES- presentado por el cliente.", "Seleccionar Consultar.",
    "Verificar identidad, evento, cantidad y total.", "Recibir el dinero.", "Seleccionar el método recibido.",
    "Pulsar Confirmar pago y emitir.", "Comprobar que la venta cambió a Pagada y que aparecieron códigos ENT-.",
])
add_note(doc, "No duplicar", "Una compra web que ya figura como Pagada no debe confirmarse nuevamente. El vendedor puede consultarla o imprimirla, pero no volver a cobrarla.", "danger")
add_heading(doc, "5.3 Historial de ventas", 2)
add_bullets(doc, ["Ventas web simuladas: vendedor mostrado como Compra en línea.", "Ventas presenciales: muestran el vendedor que recibió el pago.", "Reservas: aparecen como Pendiente hasta que se confirme el pago.", "El botón Consultar permite revisar detalles y entradas emitidas."])

add_heading(doc, "6. Funciones de Control de acceso", 1)
doc.add_paragraph("Este rol opera en la puerta del evento. Es el único responsable de marcar una entrada como utilizada.")
add_heading(doc, "6.1 Validar una entrada", 2)
add_steps(doc, ["Abrir Control de acceso.", "Escribir el código ENT- o seleccionar Escanear QR.", "Revisar evento, función, titular y estado.", "Indicar el punto de acceso.", "Seleccionar Validar ingreso.", "Confirmar que el resultado sea Aceptado."])
add_heading(doc, "6.2 Resultados posibles", 2)
add_table(doc, ["Resultado", "Significado", "Acción"], [
    ("Vigente", "Entrada pagada y no utilizada", "Autorizar y validar"),
    ("Usada", "La entrada ya fue validada", "Rechazar el nuevo intento"),
    ("Anulada", "La entrada fue invalidada", "Rechazar el ingreso"),
    ("Código inexistente", "No corresponde a una entrada", "Revisar escritura o rechazar"),
    ("Evento cancelado", "La función o evento no está disponible", "Rechazar y remitir a administración"),
])
add_note(doc, "Doble uso", "La validación se ejecuta dentro de una transacción. Aunque dos operadores consulten el mismo código, solamente la primera validación puede aceptarse.", "success")

add_heading(doc, "7. Funciones del Cliente", 1)
add_heading(doc, "7.1 Consultar el catálogo", 2)
add_steps(doc, ["Abrir la página principal.", "Revisar eventos publicados.", "Seleccionar Ver evento.", "Consultar función, precio, cupos y tipo de entrada.", "Seleccionar la opción deseada."])
add_heading(doc, "7.2 Compra web simulada", 2)
add_steps(doc, ["Iniciar sesión.", "Seleccionar cantidad y descuento opcional.", "Revisar el resumen.", "Elegir Pagar ahora (simulado).", "Confirmar la operación.", "Abrir Mis entradas y revisar los códigos ENT- generados."])
add_note(doc, "Simulación", "Este método no cobra una tarjeta real. Representa la aprobación de una pasarela y genera las entradas inmediatamente.", "info")
add_heading(doc, "7.3 Reserva para pagar en ventanilla", 2)
add_steps(doc, ["Seleccionar cantidad y revisar el resumen.", "Elegir Reservar y pagar en ventanilla.", "Anotar el código RES-.", "Abrir Mis entradas y comprobar Pendiente de pago.", "Presentar el código al vendedor.", "Pagar en ventanilla.", "Actualizar Mis entradas y comprobar los nuevos códigos ENT-."])
add_heading(doc, "7.4 Mis entradas", 2)
add_bullets(doc, ["Reservas pendientes: muestran RES-, cantidad y total por pagar.", "Entradas emitidas: muestran código ENT-, evento, función y estado.", "Una entrada Vigente puede presentarse en el ingreso.", "Una entrada Usada no permite un segundo acceso."])

add_heading(doc, "8. Flujos completos de venta", 1)
add_heading(doc, "8.1 Compra web simulada", 2)
doc.add_paragraph("Cliente → selecciona entradas → pago simulado aprobado → venta Pagada → entradas ENT- Vigente → Control de acceso → entrada Usada.")
add_heading(doc, "8.2 Reserva y pago en ventanilla", 2)
doc.add_paragraph("Cliente → crea reserva RES- → venta Pendiente y cupo reservado → vendedor recibe pago → confirma → venta Pagada → entradas ENT- Vigente → Control de acceso → entrada Usada.")
add_heading(doc, "8.3 Venta presencial", 2)
doc.add_paragraph("Cliente llega a ventanilla → vendedor identifica o registra cliente → recibe pago → Vender y emitir → venta Pagada → entradas ENT- Vigente → Control de acceso → entrada Usada.")
add_table(doc, ["Canal", "Estado inicial", "Quién confirma el pago", "Cuándo se emiten entradas"], [
    ("Web simulada", "Pagada", "Sistema simulado", "Inmediatamente"),
    ("Reserva web", "Pendiente", "Vendedor", "Después de recibir el pago"),
    ("Presencial", "Pagada", "Vendedor", "Al completar la venta"),
])

add_heading(doc, "9. Estados del sistema", 1)
add_heading(doc, "9.1 Estados de venta", 2)
add_table(doc, ["Estado", "Descripción"], [("Pendiente", "Reserva creada, todavía sin pago ni entradas."), ("Pagada", "Pago confirmado y entradas emitidas."), ("Cancelada", "Operación cancelada."), ("Reembolsada", "Pago devuelto o revertido.")])
add_heading(doc, "9.2 Estados de entrada", 2)
add_table(doc, ["Estado", "Descripción"], [("Vigente", "Puede utilizarse para ingresar."), ("Usada", "Ya fue validada en un acceso."), ("Anulada", "No permite ingreso.")])
add_heading(doc, "9.3 Identificadores", 2)
add_table(doc, ["Prefijo", "Uso", "Ejemplo"], [("VTA-", "Venta pagada", "VTA-20260720-AB12CD34"), ("RES-", "Reserva pendiente", "RES-20260720-AB12CD34"), ("ENT-", "Entrada individual", "ENT-0CAE05334A39")])

add_heading(doc, "10. Solución de problemas", 1)
add_table(doc, ["Situación", "Causa probable", "Solución"], [
    ("No aparece Seleccionar", "Venta cerrada, función iniciada o cupo agotado", "Revisar fechas, estado y cupos"),
    ("Código de descuento inválido", "Vigencia, mínimo, límite o código incorrecto", "Revisar condiciones; no utilizar ENT-"),
    ("Mis entradas está vacío", "No hay ventas pagadas o se creó una reserva", "Revisar Reservas pendientes o completar el pago"),
    ("Reserva sin entradas", "Comportamiento esperado antes del pago", "Pagar en ventanilla y pedir confirmación al vendedor"),
    ("Venta web no muestra vendedor", "Fue creada directamente por el cliente", "Se identifica como Compra en línea"),
    ("Entrada rechazada como usada", "Ya fue validada anteriormente", "Revisar fecha e historial; no autorizar de nuevo"),
    ("No abre la cámara", "Permiso denegado o navegador incompatible", "Usar Chrome/Edge en localhost o escribir el código"),
    ("Cuenta bloqueada", "Administrador bloqueó o desactivó el usuario", "Solicitar desbloqueo al administrador"),
])
add_note(doc, "Base de datos", "No elimine registros directamente en producción. Las ventas, detalles, entradas y controles están relacionados por claves foráneas; utilice siempre las funciones del sistema y realice respaldos.", "danger")

add_heading(doc, "11. Seguridad y buenas prácticas", 1)
add_bullets(doc, [
    "No compartir contraseñas entre miembros del equipo.", "Asignar únicamente el rol necesario.",
    "Bloquear inmediatamente cuentas que ya no deban ingresar.", "No confirmar una reserva antes de recibir o verificar el pago.",
    "No validar una entrada desde el módulo del vendedor.", "No publicar datos sensibles en GitHub ni incluir el archivo .env.",
    "Realizar copias de seguridad antes de modificar datos en caliente.", "Cerrar sesión al terminar, especialmente en equipos compartidos.",
])

add_heading(doc, "12. Glosario y listas de verificación", 1)
add_heading(doc, "12.1 Glosario", 2)
add_table(doc, ["Término", "Definición"], [
    ("Evento", "Actividad general publicada en el catálogo."), ("Función", "Fecha, hora y lugar específicos de un evento."),
    ("Aforo", "Cantidad máxima de asistentes permitida."), ("Tipo de entrada", "Localidad o categoría comercial, por ejemplo General o VIP."),
    ("Reserva", "Solicitud pendiente de pago que conserva cupo."), ("Entrada", "Boleto individual emitido después del pago."),
    ("Control de acceso", "Proceso que comprueba y consume una entrada vigente."),
])
add_heading(doc, "12.2 Lista previa a publicar un evento", 2)
add_bullets(doc, ["Evento con nombre, descripción e imagen.", "Organizador asignado.", "Función con fechas correctas.", "Ventana de venta vigente.", "Recinto, ciudad y aforo configurados.", "Tipos de entrada, precios y cupos creados.", "Suma de cupos menor o igual al aforo."])
add_heading(doc, "12.3 Lista del vendedor", 2)
add_bullets(doc, ["Identidad o correo verificados.", "Evento y función correctos.", "Cantidad y total revisados.", "Pago recibido o transferencia verificada.", "Venta o reserva confirmada una sola vez.", "Entradas emitidas o impresas."])
add_heading(doc, "12.4 Lista del operador de acceso", 2)
add_bullets(doc, ["Código ENT- legible.", "Evento y función correctos.", "Estado Vigente.", "Punto de acceso registrado.", "Resultado Aceptado antes de permitir el ingreso."])

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(120)
r = p.add_run("FIN DEL MANUAL")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor.from_string(PURPLE)
p = doc.add_paragraph("TicketFlow · Sistema de venta y gestión de entradas")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
